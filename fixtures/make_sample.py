# -*- coding: utf-8 -*-
"""生成测试课件 sample.pptx：封面（深色背景+标题）/ 内容页（文字+色块+表格）/ 结束页。"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, 'sample.pptx')

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank = prs.slide_layouts[6]

# ---------- P1 封面：深色背景 + 标题 ----------
s1 = prs.slides.add_slide(blank)
s1.background.fill.solid()
s1.background.fill.fore_color.rgb = RGBColor(0x1B, 0x28, 0x38)
tb = s1.shapes.add_textbox(Inches(1.2), Inches(2.4), Inches(10.9), Inches(1.4))
tf = tb.text_frame
p = tf.paragraphs[0]
p.text = '叙事理论与实践'
p.font.size = Pt(40)
p.font.bold = True
p.font.color.rgb = RGBColor(0xF7, 0xF3, 0xED)
p.alignment = PP_ALIGN.CENTER
p2 = tf.add_paragraph()
p2.text = '第一讲 · 导论'
p2.font.size = Pt(20)
p2.font.color.rgb = RGBColor(0xE8, 0x82, 0x5A)
p2.alignment = PP_ALIGN.CENTER

# ---------- P2 内容页：标题 + 正文 + 色块 + 表格 ----------
s2 = prs.slides.add_slide(blank)
tb2 = s2.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(8), Inches(1.0))
tb2.text_frame.text = '一、什么是叙事'
tb2.text_frame.paragraphs[0].runs[0].font.size = Pt(28)
tb2.text_frame.paragraphs[0].runs[0].font.bold = True

body = s2.shapes.add_textbox(Inches(0.8), Inches(1.7), Inches(6.5), Inches(4.0))
btf = body.text_frame
btf.text = '叙事是人类理解世界的底层方式。'
btf.paragraphs[0].runs[0].font.size = Pt(16)
for line in ['时间、因果与视角是叙事的三个轴。', '同一事件，不同叙述者给出不同故事。']:
    pp = btf.add_paragraph()
    pp.text = line
    pp.runs[0].font.size = Pt(16)

box = s2.shapes.add_shape(1, Inches(7.8), Inches(1.7), Inches(4.6), Inches(1.2))
box.fill.solid()
box.fill.fore_color.rgb = RGBColor(0xE8, 0x82, 0x5A)
box.line.fill.background()
bt = box.text_frame
bt.text = '重点：叙事视角'
bt.paragraphs[0].runs[0].font.size = Pt(18)
bt.paragraphs[0].runs[0].font.bold = True
bt.paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
bt.paragraphs[0].alignment = PP_ALIGN.CENTER

rows, cols = 3, 3
tbl_shape = s2.shapes.add_table(rows, cols, Inches(7.8), Inches(3.4), Inches(4.6), Inches(2.2))
tbl = tbl_shape.table
data = [['视角', '例子', '效果'], ['全知', '说书人', '俯瞰全局'], ['限知', '主人公', '身临其境']]
for r in range(rows):
    for c in range(cols):
        cell = tbl.cell(r, c)
        cell.text = data[r][c]
        cell.text_frame.paragraphs[0].runs[0].font.size = Pt(13)

# ---------- P3 结束页 ----------
s3 = prs.slides.add_slide(blank)
tb3 = s3.shapes.add_textbox(Inches(1.2), Inches(2.6), Inches(10.9), Inches(1.2))
tf3 = tb3.text_frame
tf3.text = '下次课：故事与话语'
tf3.paragraphs[0].runs[0].font.size = Pt(32)
tf3.paragraphs[0].runs[0].font.bold = True
tf3.paragraphs[0].alignment = PP_ALIGN.CENTER

prs.save(OUT)
print('written:', OUT, '%.1f KB' % (os.path.getsize(OUT) / 1024))
